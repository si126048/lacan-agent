from __future__ import annotations
from pathlib import Path


def extract_office_pages(path: Path) -> tuple[list[str], dict]:
    suffix = path.suffix.lower()
    if suffix == '.docx':
        return _extract_docx(path)
    if suffix == '.xlsx':
        return _extract_xlsx(path)
    if suffix == '.pptx':
        return _extract_pptx(path)
    raise ValueError(f'unsupported office format: {suffix}')


def _extract_docx(path: Path) -> tuple[list[str], dict]:
    try:
        from docx import Document
    except ImportError:
        raise ImportError('python-docx is required: pip install python-docx')
    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        parts.append(para.text)
    text = '\n\n'.join(p for p in parts if p.strip())
    return [text], {'format': 'docx', 'page_count': 1}


def _extract_xlsx(path: Path) -> tuple[list[str], dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError('openpyxl is required: pip install openpyxl')
    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else '' for c in row]
            rows.append('\t'.join(cells))
        if any(rows):
            parts.append(f'## Sheet: {sheet_name}\n' + '\n'.join(rows))
    wb.close()
    text = '\n\n'.join(parts)
    return [text], {'format': 'xlsx', 'page_count': 1, 'sheets': len(wb.sheetnames)}


def _extract_pptx(path: Path) -> tuple[list[str], dict]:
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError('python-pptx is required: pip install python-pptx')
    prs = Presentation(str(path))
    pages: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            pages.append(f'--- Slide {i} ---\n' + '\n'.join(texts))
    return pages, {'format': 'pptx', 'page_count': len(pages)}
